from docx import Document
from docx.shared import Pt

doc = Document()

# Titre
doc.add_heading("Carnet de projet : Similarité de joueurs de football", level=0)

# 1. Objectif
doc.add_heading("1. Objectif", level=1)
doc.add_paragraph(
    "Identifier, pour un joueur de référence, les joueurs aux profils "
    "statistiques les plus proches, puis visualiser la comparaison entre "
    "le joueur de référence et son voisin le plus proche sous forme de "
    "graphique radar."
)

# 2. Structure du projet
doc.add_heading("2. Structure du projet", level=1)
for item in [
    "main.py : script principal (chargement, calcul, affichage)",
    "similarity.py : fonctions charger_donnees, standardiser, joueurs_similaires",
    "visualisation.py : fonction radar_comparaison",
    "data/players.csv : base de données des joueurs",
    "radar.png : graphique radar généré",
]:
    doc.add_paragraph(item, style="List Bullet")

# 3. Détail du code
doc.add_heading("3. Détail du code", level=1)

doc.add_heading("3.1 Importations", level=2)
doc.add_paragraph(
    "matplotlib.pyplot sert à afficher la figure (plt.show()). "
    "Les trois fonctions de similarity.py gèrent les données et le calcul "
    "de similarité. radar_comparaison gère la visualisation."
)
doc.add_paragraph(
    "import matplotlib.pyplot as plt\n"
    "from similarity import charger_donnees, standardiser, joueurs_similaires\n"
    "from visualisation import radar_comparaison"
).style = doc.styles["No Spacing"]

doc.add_heading("3.2 Paramètres", level=2)
doc.add_paragraph(
    "Ces deux constantes centralisent les réglages : le chemin du fichier "
    "et le joueur de référence. Les modifier suffit pour changer l'analyse."
)
doc.add_paragraph(
    'CHEMIN_DONNEES = "data/players.csv"\n'
    'NOM_JOUEUR = "Joueur 05"'
).style = doc.styles["No Spacing"]

doc.add_heading("3.3 Chargement des données", level=2)
doc.add_paragraph(
    "Lit le fichier CSV et le transforme en DataFrame pandas : une ligne "
    "par joueur, une colonne par statistique."
)
doc.add_paragraph(
    "df = charger_donnees(CHEMIN_DONNEES)"
).style = doc.styles["No Spacing"]

doc.add_heading("3.4 Standardisation", level=2)
doc.add_paragraph(
    "Pour chaque statistique numérique, on calcule le z-score : "
    "z = (x − moyenne) / écart-type. Toutes les statistiques ont ainsi "
    "une moyenne de 0 et un écart-type de 1, ce qui évite qu'une grande "
    "valeur brute (ex. passes) domine les autres. La fonction renvoie le "
    "DataFrame enrichi des colonnes _z et la liste z_cols."
)
doc.add_paragraph(
    "df, z_cols = standardiser(df)"
).style = doc.styles["No Spacing"]

doc.add_heading("3.5 Recherche des joueurs similaires", level=2)
doc.add_paragraph(
    "Pour le joueur de référence, on calcule la distance euclidienne avec "
    "chaque autre joueur sur les colonnes z_cols : "
    "distance = √ Σ (z_joueur − z_référence)². On trie par distance croissante "
    "et on garde les n = 5 premiers. Une distance faible signifie un profil proche."
)
doc.add_paragraph(
    'resultats = joueurs_similaires(df, z_cols, NOM_JOUEUR, n=5)\n'
    'print(f"Joueurs les plus proches de {NOM_JOUEUR} :")\n'
    'print(resultats)'
).style = doc.styles["No Spacing"]

doc.add_heading("3.6 Sélection du voisin le plus proche", level=2)
doc.add_paragraph(
    "iloc[0] prend la première ligne du classement, donc le joueur le plus "
    "proche. ['name'] extrait son nom."
)
doc.add_paragraph(
    'plus_proche = resultats.iloc[0]["name"]'
).style = doc.styles["No Spacing"]

doc.add_heading("3.7 Radar de comparaison", level=2)
doc.add_paragraph(
    "radar_comparaison trace un graphique radar : chaque axe correspond à une "
    "statistique standardisée, et chaque joueur est un polygone. "
    "chemin_sortie sauvegarde l'image sur le disque ; plt.show() l'affiche."
)
doc.add_paragraph(
    'fig = radar_comparaison(df, z_cols, NOM_JOUEUR, plus_proche,\n'
    '                        chemin_sortie="radar.png")\n'
    'plt.show()'
).style = doc.styles["No Spacing"]

# 4. Points de vigilance
doc.add_heading("4. Points de vigilance", level=1)
for item in [
    "Vérifier que NOM_JOUEUR existe exactement dans la colonne name.",
    "Exclure le joueur de référence de son propre classement (sinon distance 0).",
    "Si toutes les distances valent 0.0, vérifier la création des colonnes en amont.",
]:
    doc.add_paragraph(item, style="List Bullet")

# 5. Résultats attendus
doc.add_heading("5. Résultats attendus", level=1)
for item in [
    "Un tableau de 5 joueurs classés par similarité dans le terminal.",
    "Un graphique radar affiché et sauvegardé dans radar.png.",
]:
    doc.add_paragraph(item, style="List Bullet")

# 6. Enseignements
doc.add_heading("6. Enseignements", level=1)
for item in [
    "Découper le projet en modules (données, calcul, visualisation) rend le code lisible et réutilisable.",
    "Standardiser avant de comparer est indispensable quand les échelles diffèrent.",
    "Afficher les résultats intermédiaires aide à détecter rapidement les erreurs.",
]:
    doc.add_paragraph(item, style="List Bullet")

doc.save("carnet_projet.docx")
print("Document créé : carnet_projet.docx")